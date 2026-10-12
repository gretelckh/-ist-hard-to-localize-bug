#!/usr/bin/env python3
"""Add an '错误原因分析' (root-cause analysis) column to all_bugs_analysis.xlsx
and reconcile column 17 (异常类型) with column 19 (故障范围标记 / three-channel).

For each of the 49 REPRODUCED bugs, this writes a plain-language root-cause analysis
in the SAME style used in the paper discussion: what the SDK actually reported,
where the real defect lies, and why it is hard to localise / fix.

Run from MCP-Fault-Benchmark/:
    python analysis/backfill_root_cause.py
"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool'))
XL = os.path.join(BASE, 'all_bugs_analysis.xlsx')

# issue -> (root-cause analysis in Chinese, reconciled anomaly-type label)
ROOT_CAUSE = {
  # --- pushed to Server ---
  110: ("Server 返回非 UTF-8 字节，客户端 JSON 解码失败，但真实根因（字节解码错误）被 SDK 通用化成 "
        "`Error executing tool bad` 在 Server 工具分派层报出。源是 Client 层，报告层却是 Server，属层错位。",
        'pushed-server'),
  1361: ("Client 层 `call_tool` 抛出的异常被 Server 吞掉并重报为 `Error executing tool boom`（is_error=True），"
         "真实异常值丢失。源是 Client，报告层是 Server，属层错位 + 通用化掩蔽。", 'pushed-server'),
  2114: ("`call_tool` 内部异常被 `ExceptionGroup` 包裹，展开后仍只剩 `Error executing tool boom`，"
         "真实根因被通用化掩蔽；源是 Client，报告层是 Server。", 'pushed-server'),
  3637: ("Host 传扁平参数 dict（期望 Pydantic model 被 flatten），但 Server 侧未 flatten，按缺失字段报校验错误"
         "`Error executing tool summarize`（is_error=True）。源是 Host，报告层是 Server，属层错位。", 'pushed-server'),
  3572: ("Host 传错参数类型（数字而非字符串），Server 校验错误回显了被拒绝的值 `Error executing tool lookup`。"
         "源是 Host，报告层是 Server，属层错位（校验错误被推到 Server 层）。", 'pushed-server'),
  3579: ("Host 向一个无组件的 server 请求未知工具，报 `Unknown tool`（is_error=True）。"
         "源是 Host 的请求，报告层是 Server，属层错位。", 'pushed-server'),
  2741: ("有状态 Streamable HTTP 会话内的异常被掩蔽为通用 `Error executing tool boom`（is_error=True），"
         "真实 cause 只在独立的 `Session crashed` 日志中。源是 Transport，报告层是 Server，属层错位 + 掩蔽。", 'pushed-server'),
  # --- silent ---
  2057: ("`id:null` 的 JSON-RPC 消息被当作 notification 静默接受（HTTP 202），未按规范拒绝。"
         "逻辑错误无任何报错标志，属静默失效。", 'silent'),
  3067: ("`call_tool(add,{a:1,b:2,typo:9})` 返回 3 无报错，拼错的额外参数被静默丢弃（extra=ignore）。"
         "输入错误被无声吞掉，属静默失效。", 'silent'),
  2539: ("可选的 `None` 字段被序列化为 JSON `null`，返回 `{\"text\": null}`，is_error=False。"
         "错误的结果结构被当作成功返回，属静默失效。", 'silent'),
  2618: ("`MCP-Protocol-Version` header 与 body 的 `protocolVersion` 不一致，但服务端按 body 协商返回 HTTP 200，"
         "未校验 header。协议错配被静默接受，属静默失效。", 'silent'),
  3549: ("20 个并发调用共享状态导致丢失更新（final=1，应=20），is_error=False。"
         "并发竞态结果错误但无报错，属静默失效。", 'silent'),
  3578: ("同名工具重复注册时静默保留第一个版本（shadow），第二个被忽略，返回 2 无报错。"
         "注册冲突被无声吞掉，属静默失效。", 'silent'),
  3585: ("`format=` 输出未注册的 MIME 类型 `image/jpg`，is_error=False。"
         "非规范媒体类型被静默接受，属静默失效。", 'silent'),
  3592: ("服务端返回 stale-cache-fallback 标记作为成功结果（is_error 与真实状态不匹配），is_error=False。"
         "状态错配被当作成功，属静默失效。", 'silent'),
  3640: ("未知的 ServerCapabilities key 在广播能力时被静默丢弃（is_error=False）。"
         "能力信息丢失但无报错，属静默失效。", 'silent'),
  3643: ("prompt handler 返回 UserList 被渲染成单条字符串化消息（is_error=False）。"
         "结果结构错误被当作成功，属静默失效。", 'silent'),
  3646: ("Host 传入的未知额外工具参数被静默接受（无 opt-in 拒绝），is_error=False。"
         "参数错误被无声吞掉，属静默失效。", 'silent'),
  # --- caller-exception ---
  3556: ("客户端把 404（死会话）当作终结而非重新初始化，抛出 `ExceptionGroup`（裸异常）。"
         "Transport 层错误以调用方裸异常暴露，属调用方裸抛通道。", 'caller-exception'),
  3557: ("畸形的非 JSON-RPC body 被回答为 `-32602`（应 `-32600`），HTTP 400 裸异常。"
         "错误码规范问题，以调用方裸异常暴露。", 'caller-exception'),
  3560: ("慢工具无默认超时，阻塞调用方，抛出 `ExceptionGroup`。"
         "Transport 层超时以调用方裸异常暴露，属调用方裸抛通道。", 'caller-exception'),
  # --- single-layer-correct ---
  234: ("工具不存在时报错类型不精准：期望 `is_error=False` 却返回 True，属错误分类粒度问题。"
         "报在 Server 层（正确），但异常类型语义不准。", 'single-layer'),
  1579: ("`read_resource` 找不到资源时返回错误码不对（应是特定 JSON-RPC code，实际返回其他）。"
         "报在 Server 层（正确），但错误码不符规范。", 'single-layer'),
  2422: ("工具不存在时应抛专用 `ToolNotFoundError`，实际抛笼统 `ToolError`，错误分类粒度不足。"
         "报在 Server 层（正确），但异常类型不精准。", 'single-layer'),
  3341: ("工具返回非有限浮点数 `nan`，结构化内容校验失败（`None is not of type number`）。"
         "报在 Server 层（正确），是输出 schema 校验问题。", 'single-layer'),
  3342: ("prompt 校验错误被压成笼统 `MCPError: Internal server error`，真实校验细节丢失。"
         "报在 Server 层（正确），但错误信息被简化。", 'single-layer'),
  3351: ("输入校验时丢弃 jsonschema 字段信息，报 `Error executing tool f: 1 validation error`。"
         "报在 Server 层（正确），但校验细节不完整。", 'single-layer'),
  3385: ("工具返回非有限浮点数 `inf` 导致 text content 无效 JSON，结构化校验失败。"
         "报在 Server 层（正确），是输出序列化/校验问题。", 'single-layer'),
  3644: ("outputSchema 处于 validation 模式但内容按 alias 序列化，导致 schema 不匹配，抛 RuntimeError。"
         "报在 Server 层（正确），是输出 schema 校验问题。", 'single-layer'),
  # --- single-layer-masked ---
  348: ("任意工具结果无法置 isError=True；真实错误（如 PydanticSchemaGenerationError）被通用化成 "
        "`Error executing tool img`。报在 Server 层，但真实根因被掩蔽。", 'single-layer'),
  417: ("服务端内部错误被通用化成 `Error executing tool boom`，真实异常值丢失。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  698: ("工具抛异常后真实异常值被吞，只剩 `Error executing tool boom`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  743: ("工具根因未被正确记录，报 `Error executing tool boom`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  891: ("`dict+str` 非法操作的真实 TypeError 被通用化成 `Error executing tool f`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  1060: ("Image 类型触发 PydanticSchemaGenerationError 被通用化成 `Error executing tool img`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  1702: ("工具输出未过 outputSchema 校验，被通用化成 `Error executing tool get_name`，隐藏 schema 细节。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  2153: ("handler 应抛异常但异常值被吞，报 `Error executing tool boom`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  2386: ("CWE-209 敏感信息（password=abc）被清洗，真实 ValueError 被通用化成 `Error executing tool boom`。"
         "报在 Server 层，但敏感根因被掩蔽。", 'single-layer'),
  2431: ("McpError 非 pickle-safe，unpickle 失败被通用化成 `Error executing tool boom`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  2770: ("McpError 被吞成 isError:true，报 `Error executing tool boom`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3099: ("Field(alias) 在 schema 中但非法 Python 关键字，报 `Error executing tool ext`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3266: ("工具调用中的异常未记录，报 `Error executing tool boom`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3656: ("`Audio(data=b\"\")` 抛异常而 `Image(data=b\"\")` 正常，报 `Error executing tool a`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3554: ("非 UTF-8 字节输出崩溃，报 `Error executing tool get_bytes`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3569: ("工具抛非 ToolError，被通用化成 `Error executing tool fragile2`。报在 Server 层，但根因被掩蔽。", 'single-layer'),
  3586: ("handler 抛非 ToolError，客户端暴露不清晰，报 `Error executing tool fragile`。"
         "报在 Server 层，但根因被掩蔽。", 'single-layer'),
  # --- registration-error ---
  244: ("`resource()` 不接受 `ctx: Context`，server 注册时抛 ValueError（URI 模板变量缺失）。"
         "注册期错误：server 无法构建，属另一类难定位根因。", 'registration'),
  1974: ("`find_context_parameter` 对 callable class 失败，注册时抛 AttributeError（Handler 无 __name__）。"
         "注册期错误：server 无法构建。", 'registration'),
  3573: ("`-> Iterator` 注解触发 PydanticSchemaGenerationError，注册时失败。"
         "注册期错误：server 无法构建。", 'registration'),
}

def main():
    wb = openpyxl.load_workbook(XL); ws = wb.active
    hdr = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
    # add a new column '错误原因分析' if absent
    if '错误原因分析' not in hdr:
        newcol = ws.max_column + 1
        hdr['错误原因分析'] = newcol - 1
        ws.cell(row=1, column=newcol, value='错误原因分析')
        ws.cell(row=1, column=newcol).font = openpyxl.styles.Font(bold=True, color='FFFFFF')
        ws.cell(row=1, column=newcol).fill = openpyxl.styles.PatternFill('solid', fgColor='4472C4')
        ws.column_dimensions[openpyxl.utils.get_column_letter(newcol)].width = 80
    exc_col = hdr['异常类型']
    rc_col = hdr['错误原因分析']
    n = 0
    for row in ws.iter_rows(min_row=2):
        issue_cell = row[2].value
        if not issue_cell:
            continue
        issue = int(str(issue_cell).lstrip('#'))
        if issue in ROOT_CAUSE:
            reason, chan = ROOT_CAUSE[issue]
            row[rc_col].value = reason
            # FORCE 异常类型 to the canonical three-channel label so column 17 is
            # always consistent with column 19 (故障范围标记). Old 'see-raw' / stale
            # labels are replaced; single-layer-correct & single-layer-masked both map
            # to 'single-layer', registration maps to 'registration'.
            row[exc_col].value = chan
            n += 1
    wb.save(XL)
    print('filled 错误原因分析 for', n, 'bugs ->', XL)

if __name__ == '__main__':
    main()