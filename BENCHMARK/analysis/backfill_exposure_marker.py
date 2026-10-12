#!/usr/bin/env python3
"""Backfill column '故障范围标记' (fault-scope / exposure-channel marker) in
all_bugs_analysis.xlsx from the evidence-based three-channel classification.

Source of truth: data/full_pool/active_dataset_analysis.json -> tier_b.per_bug
Each reproduced bug gets an exposure-channel label. Non-reproduced bugs keep the
honest 'not-exercised' marker (no runtime evidence => no channel can be inferred).

Run from MCP-Fault-Benchmark/:
    python analysis/backfill_exposure_marker.py
"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool'))
XL = os.path.join(BASE, 'all_bugs_analysis.xlsx')
AN = os.path.join(BASE, 'active_dataset_analysis.json')

CHANNEL_CN = {
    'pushed-server': '三通道① 推到Server(错位:源非Server被报成Server)',
    'silent': '三通道② 静默(is_error=False,逻辑错被吞)',
    'caller-exception': '三通道③ 调用方裸抛异常(Transport/HTTP)',
    'single-layer-correct': '单层-正确定位(Server源,报在Server)',
    'single-layer-masked': '单层-正确层但通用化掩蔽(Error executing tool)',
    'registration-error': '注册期错误(server构建/注册失败)',
}

def main():
    an = json.load(open(AN, encoding='utf-8'))
    per_bug = {b['issue']: b for b in an['tier_b']['per_bug']}

    wb = openpyxl.load_workbook(XL); ws = wb.active
    hdr = {c.value: i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
    col = hdr['故障范围标记']  # 0-based index
    n_updated = 0
    for row in ws.iter_rows(min_row=2):
        issue_cell = row[2].value  # 'GitHub Issue' column (#NNN)
        if not issue_cell:
            continue
        issue = int(str(issue_cell).lstrip('#'))
        b = per_bug.get(issue)
        if b:
            row[col].value = CHANNEL_CN.get(b['channel'], b['channel'])
            n_updated += 1
    wb.save(XL)
    print('backfilled 故障范围标记 for', n_updated, 'reproduced bugs ->', XL)

if __name__ == '__main__':
    main()