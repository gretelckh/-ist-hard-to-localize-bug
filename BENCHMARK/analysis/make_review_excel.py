"""Generate the final review Excel (all 439 with verdicts) + confirmed-bug list.

Run from MCP-Fault-Benchmark/:
    python analysis/make_review_excel.py
"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
POOL = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'full_pool_screened.json'))
VERDICT = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'bug_review_verdicts.json'))
OUT = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'bug_review_full.xlsx'))
CONF = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'confirmed_bugs.json'))

HEADER = ['序号', 'Issue', '标题', '初筛label', '人工复核verdict', '原因', 'dup_of', 'body前300']


def main():
    pool = json.load(open(POOL, encoding='utf-8'))
    verdicts = json.load(open(VERDICT, encoding='utf-8'))['verdicts']
    bugs = [x for x in pool if x['label'] == 'BUG']
    bugs.sort(key=lambda x: x['number'])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'review'
    ws.append(HEADER)
    for i, x in enumerate(bugs, 1):
        v = verdicts.get(str(x['number']) if str(x['number']) in verdicts else x['number'], {})
        ws.append([
            i, x['number'], x['title'], 'BUG', v.get('verdict', ''), v.get('reason', ''),
            v.get('dup_of', ''), (x.get('body_head') or '')[:300],
        ])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    wb.save(OUT)

    # confirmed bugs json
    confirmed = [x for x in bugs if verdicts.get(str(x['number']) if str(x['number']) in verdicts else x['number'], {}).get('verdict') == 'BUG']
    json.dump([{'number': x['number'], 'title': x['title'], 'url': x['url']} for x in confirmed],
              open(CONF, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    print('review excel ->', OUT, '(', len(bugs), 'rows )')
    print('confirmed bugs ->', CONF, '(', len(confirmed), ')')


if __name__ == '__main__':
    main()