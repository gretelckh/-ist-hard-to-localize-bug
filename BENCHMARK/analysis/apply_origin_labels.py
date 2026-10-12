#!/usr/bin/env python3
"""Apply the authoritative expert origin-layer labels (Step 2 final output).

Merges three sources of truth for each of the 398 confirmed bugs:
  1. EXPERT_ORIGIN   (from author_origin_labels.py) - expert-read label for ALL 398
  2. verified origin (from data/real_summary.json)  - empirical origin for the 18 known-
                        reproducible bugs (OVERRIDES expert label, since empirically grounded)
  3. basic fields    (confirmed_bugs.json + full_pool_screened.json) - id/title/url/state/created/body

Produces:
  data/full_pool/origin_labels.json   authoritative per-bug origin + source
  data/full_pool/origin_labels.xlsx   Excel of the authoritative origin per bug
  (console) distribution summary

Run from MCP-Fault-Benchmark/:
    python analysis/apply_origin_labels.py
"""
import json, os, sys
sys.stdout.reconfigure(encoding='utf-8')
from collections import Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import importlib.util

HERE = os.path.dirname(os.path.abspath(__file__))
CONF = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'confirmed_bugs.json'))
POOL = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'full_pool_screened.json'))
REALSUM = os.path.normpath(os.path.join(HERE, '..', 'data', 'real_summary.json'))
OUT = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'origin_labels.json'))
OUTX = os.path.normpath(os.path.join(HERE, '..', 'data', 'full_pool', 'origin_labels.xlsx'))

# load expert labels
spec = importlib.util.spec_from_file_location('aol', os.path.join(HERE, 'author_origin_labels.py'))
aol = importlib.util.module_from_spec(spec); spec.loader.exec_module(aol)
EXPERT = aol.EXPERT_ORIGIN


# issues the author flagged as ambiguous / needing re-review. These rows are highlighted
# in the Excel so the author can re-verify the origin-layer judgement.
REVIEW = {
    201: ('Host-or-Client', 'pure-Python client hangs on ctx.info() logging; root cause could be host app or client session'),
    788: ('Client-or-Host', 'ClientSessionGroup aggregation; closing-multi-server error (client vs host boundary)'),
    915: ('Client-or-Host', 'ClientSessionGroup exception when streamable_http server unavailable'),
    3384: ('Client-or-Host', 'ClientSessionGroup KeyError when server exposes no components'),
    3490: ('Client-or-Host', 'ClientSessionGroup rejected connect leaves transport running'),
}


def main():
    conf = json.load(open(CONF, encoding='utf-8'))
    pool = json.load(open(POOL, encoding='utf-8'))
    pool_map = {x['number']: x for x in pool}
    bynum = {x['number']: x for x in conf}

    # verified origins from real reproduction (empirical, overrides expert)
    verified = {}
    if os.path.exists(REALSUM):
        rs = json.load(open(REALSUM, encoding='utf-8'))
        for c in rs.get('per_case', []):
            n = int(c['issue_src'].split('#')[1])
            verified[n] = c['origin']

    rows = []
    for n in sorted(EXPERT):
        rec = bynum.get(n)
        if not rec:
            continue
        layer = EXPERT[n]
        src = 'expert'
        if n in verified:
            if verified[n] != layer:
                layer = verified[n]
            src = 'verified'
        body = (pool_map.get(n, {}).get('body') or '') or (rec.get('body_head') or '')
        review_reason = None
        if n in REVIEW:
            review_reason = REVIEW[n][1]
        rows.append({
            'number': n, 'title': rec.get('title') or '', 'state': rec.get('state') or '',
            'created_at': rec.get('created_at') or '', 'url': rec.get('url') or '',
            'origin': layer, 'origin_source': src, 'review_reason': review_reason,
            'body': body[:400],
        })

    dist = Counter(r['origin'] for r in rows)
    print('=' * 72)
    print('AUTHORITATIVE ORIGIN LAYER (expert + verified) for %d confirmed bugs' % len(rows))
    print('=' * 72)
    for l in ['Host', 'Client', 'Server', 'Transport']:
        print('  %-12s %3d  (%s)' % (l, dist.get(l, 0), '%.1f%%' % (dist.get(l, 0) / len(rows) * 100)))
    print('  verified =', sum(1 for r in rows if r['origin_source'] == 'verified'),
          '| expert =', sum(1 for r in rows if r['origin_source'] == 'expert'))
    print('  flagged-for-review =', sum(1 for r in rows if r.get('review_reason')))

    json.dump(rows, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print('\n->', OUT)

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'origin_labels'
    headers = ['序号', 'Case ID', 'GitHub Issue', '链接', '标题', '状态', '创建时间',
               '原始层(origin)', '标签来源', '待复核', 'body(前400字)']
    ws.append(headers)
    for h in headers:
        c = ws[f'{chr(65+headers.index(h))}1']
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4472C4')
        c.alignment = Alignment(horizontal='center', vertical='center')
    red_fill = PatternFill('solid', fgColor='FFC7CE')
    red_font = Font(color='9C0006')
    for i, r in enumerate(rows, start=2):
        ws.append([i-1, 'BUG-%04d' % r['number'], '#' + str(r['number']), r['url'],
                   r['title'], r['state'], r['created_at'],
                   r['origin'], r['origin_source'],
                   ('Y' if r.get('review_reason') else ''),
                   r['body']])
        if r.get('review_reason'):
            for col in range(1, 12):
                cell = ws.cell(row=i, column=col)
                cell.fill = red_fill
                cell.font = red_font
    widths = [6, 12, 12, 50, 55, 8, 20, 16, 10, 8, 60]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64+j)].width = w
    wb.save(OUTX)
    print('->', OUTX)


if __name__ == '__main__':
    main()